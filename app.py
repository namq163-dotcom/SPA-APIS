import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from datetime import datetime
from io import BytesIO, StringIO
from openpyxl import load_workbook

# Cấu hình giao diện
st.set_page_config(page_title="Sun PhuQuoc Airways - APIS", page_icon="☀️", layout="wide")

# ==========================================
# CÁC HÀM XỬ LÝ DỮ LIỆU
# ==========================================
def parse_date(date_str):
    if pd.isna(date_str) or not str(date_str).strip() or str(date_str).strip().lower() == 'nan': return ""
    date_str = str(date_str).strip()
    try:
        d = datetime.strptime(date_str, "%d/%m/%y")
        if d.year > 2050: d = d.replace(year=d.year - 100)
        return d.strftime("%d/%m/%Y")
    except: return date_str

def process_roster_data_vn(gd_file, template_file_path):
    content = gd_file.getvalue()
    df_gd = None
    
    # Thử các engine đọc file
    for engine in ['openpyxl', 'xlrd']:
        try: 
            df_gd = pd.read_excel(BytesIO(content), engine=engine)
            if df_gd is not None and len(df_gd) > 0: break
        except: continue

    if df_gd is None or len(df_gd) == 0:
        raise ValueError("Không thể đọc được dữ liệu. Định dạng file không hỗ trợ.")

    # 1. TỰ ĐỘNG TÌM NƠI ĐI / NƠI ĐẾN
    dep_place, arr_place = "N/A", "N/A"
    # Quét 15 dòng đầu để tìm thông tin hành trình
    search_area = df_gd.iloc[:15].fillna("").astype(str).values
    for row in search_area:
        row_str = " ".join(row).lower()
        if "from" in row_str:
            dep_place = row_str.split("from")[-1].split("-")[0].strip().upper()
        if "to" in row_str:
            arr_place = row_str.split("to")[-1].split("-")[0].strip().upper()

    # 2. TÌM HEADER BẢNG TỔ BAY
    header_idx = None
    for idx, row in df_gd.iterrows():
        row_str = row.fillna("").astype(str).str.lower()
        if any('passport' in s for s in row_str) and any('name' in s for s in row_str):
            header_idx = idx
            break
            
    if header_idx is None: raise ValueError("Không tìm thấy bảng danh sách tổ bay.")
        
    header_row = df_gd.iloc[header_idx].fillna("").astype(str).str.lower()
    col_name = next((i for i, v in enumerate(header_row) if 'name' in v), None)
    col_passport = next((i for i, v in enumerate(header_row) if 'passport' in v and 'expiry' not in v), None)
    col_dob = next((i for i, v in enumerate(header_row) if 'birth' in v), None)
    col_gender = next((i for i, v in enumerate(header_row) if v == 'g'), None)
    col_nat = next((i for i, v in enumerate(header_row) if 'ntly' in v), None)
    col_expiry = next((i for i, v in enumerate(header_row) if 'expiry' in v), None)

    crew_data = []
    seen_passports = set()
    
    for idx in range(header_idx + 1, len(df_gd)):
        row = df_gd.iloc[idx]
        if str(row.iloc[0]).lower() == 'nan': break
        
        passport_val = str(row.iloc[col_passport]).strip() if col_passport is not None else ""
        if passport_val in seen_passports or passport_val == 'nan': continue
        seen_passports.add(passport_val)
        
        name_val = str(row.iloc[col_name]).strip()
        name_parts = name_val.split()
        family_name = name_parts[0] if name_parts else ""
        given_name = " ".join(name_parts[1:]) if len(name_parts) > 1 else ""
        
        crew_data.append([
            family_name, None, given_name, str(row.iloc[col_gender]), str(row.iloc[col_nat]), 
            parse_date(row.iloc[col_dob]), 'P', passport_val, str(row.iloc[col_nat]), 
            parse_date(row.iloc[col_expiry])
        ])

    # 3. GHI VÀO TEMPLATE
    output = BytesIO()
    book = load_workbook(template_file_path)
    sheet = book.active
    
    # Ghi thông tin hành trình (Thay đổi tọa độ C5, C6 nếu cần)
    sheet['C5'] = dep_place
    sheet['C6'] = arr_place
    
    # Ghi danh sách
    for r_idx, row_data in enumerate(crew_data, 14):
        for c_idx, value in enumerate(row_data, 1):
            sheet.cell(row=r_idx, column=c_idx, value=value)
            
    book.save(output)
    return output.getvalue(), pd.DataFrame(crew_data)

# ==========================================
# GIAO DIỆN CHÍNH
# ==========================================
st.markdown("### ☀️ Sun PhuQuoc Airways - APIS Tool")
uploaded_gd = st.file_uploader("Tải lên file GD", type=["xlsx", "xls"])

if uploaded_gd:
    try:
        excel_data, preview = process_roster_data_vn(uploaded_gd, "Template_VNAPIS.xlsx")
        st.success("Đã xử lý xong!")
        st.download_button("Tải file APIS", excel_data, "APIS_Export.xlsx")
    except Exception as e:
        st.error(f"Lỗi: {e}")
